#!/usr/bin/env python3
"""
Tally ageing engine - Debtor & Creditor ageing report + dashboard (monthly).

Input : Tally "Group Outstandings" ageing export (.xlsx) for Sundry Debtors
        and/or Sundry Creditors (Display > Statements of Accounts > Outstandings
        > Group > Sundry Debtors/Creditors, F6 Age-wise, export to Excel).
Output: <Type>_Ageing_<Mon-YY>.xlsx  (Dashboard, Ageing Detail, Credit/Debit
        balances, Movement vs prior month, Checks, Source) and
        Ageing_Dashboard_<Mon-YY>.html (interactive, self-contained).

Usage : python tally_ageing.py DEBTORS.xlsx [CREDITORS.xlsx]
              [--prev PREV_DEBTORS.xlsx [PREV_CREDITORS.xlsx]] [--out DIR]
              [--immaterial 1]
"""
import argparse, datetime as dt, html, json, os, re, sys
import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.formatting.rule import CellIsRule

NAVY = "1F3864"; BLUE = "2A78D6"; GREY = "F2F2F2"; YELLOW = "FFF2CC"
RAMP = ["86B6EF", "5598E7", "2A78D6", "1C5CAB", "104281", "0D366B", "9E9E9E"]
NUM = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'

# ----------------------------------------------------------------- parsing
def _f(v):
    if v in (None, ""): return 0.0
    if isinstance(v, (int, float)): return float(v)
    s = str(v).replace(",", "").strip()
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    try: x = float(re.sub(r"[^0-9.\-]", "", s) or 0)
    except ValueError: return 0.0
    return -x if neg else x

def _period_end(text):
    m = re.search(r"to\s+(\d{1,2}-[A-Za-z]{3}-\d{2,4})", text or "")
    if not m: return None
    for fmt in ("%d-%b-%y", "%d-%b-%Y"):
        try: return dt.datetime.strptime(m.group(1), fmt).date()
        except ValueError: pass

def parse(path):
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    rows = list(ws.iter_rows())
    hdr = next(i for i, r in enumerate(rows) if str(r[0].value or "").strip().lower() == "particulars")
    company = str(rows[0][0].value or "").strip()
    pre = [str(r[0].value or "").strip() for r in rows[:hdr]]
    kind = "Creditors" if any("creditor" in p.lower() for p in pre) else "Debtors"
    period = next((p for p in pre if re.search(r"\d{1,2}-[A-Za-z]{3}-\d{2}", p)), "")
    as_at = _period_end(period) or _period_end(str(rows[hdr][1].value))
    blab = rows[hdr + 1]; drcr = rows[hdr + 2]
    # bucket pairs: label row has label in the Debit column
    pairs = []
    for c in range(1, ws.max_column):
        lab = blab[c].value
        if lab and str(drcr[c].value or "").lower().startswith("debit"):
            cr = c + 1 if c + 1 < ws.max_column and str(drcr[c + 1].value or "").lower().startswith("credit") else None
            pairs.append((re.sub(r"[()\s]+", " ", str(lab)).strip(), c, cr))
    if not pairs: raise SystemExit("Could not find bucket header row (Debit/Credit) under 'Particulars'.")
    total_pair, bucket_pairs = pairs[0], pairs[1:]
    if "pending" not in total_pair[0].lower():   # export without a total column
        total_pair, bucket_pairs = None, pairs
    sign = 1 if kind == "Debtors" else -1
    net = lambda r, p: sign * (_f(r[p[1]].value) - (_f(r[p[2]].value) if p[2] is not None else 0))

    parties, groups, grand = [], [], None
    top = "Sundry " + kind; cur = top
    for r in rows[hdr + 3:]:
        name = str(r[0].value or "").strip()
        if not name: continue
        if name.lower().startswith("grand total"):
            grand = {"total": net(r, total_pair) if total_pair else None,
                     "buckets": [net(r, p) for p in bucket_pairs]}
            break
        c0 = r[0]
        indent = (c0.alignment.indent or 0) if c0.alignment else 0
        if c0.font and c0.font.b and indent == 0:          # sub-group subtotal row
            cur = name; groups.append(name); continue
        grp = cur if indent > 0 else top
        if indent == 0: cur = top
        b = [net(r, p) for p in bucket_pairs]
        tot = net(r, total_pair) if total_pair else sum(b)
        code = (re.match(r"^\s*(\d+)\s", name) or [None, ""])[1]
        parties.append({"name": name, "code": code, "group": grp, "buckets": b, "total": tot})
    return {"file": os.path.basename(path), "company": company, "kind": kind, "period": period,
            "as_at": as_at, "buckets": [p[0] for p in bucket_pairs], "parties": parties,
            "groups": groups, "grand": grand, "ws": ws, "hdr": hdr}

# ----------------------------------------------------------------- analysis
def analyse(d, prev=None, immaterial=1.0):
    B = d["buckets"]; n = len(B)
    oa = next((i for i, b in enumerate(B) if "account" in b.lower()), None)
    aged = [i for i in range(n) if i != oa]
    first_overdue = aged[1:] if len(aged) > 1 else []
    over180 = [i for i in aged if re.search(r"180|365", B[i])]
    for p in d["parties"]:
        p["material"] = abs(p["total"]) >= immaterial or any(abs(x) >= immaterial for x in p["buckets"])
        p["oldest"] = next((B[i] for i in reversed(aged) if p["buckets"][i] > immaterial), "-")
        p["over90"] = sum(p["buckets"][i] for i in first_overdue)
        p["over180"] = sum(p["buckets"][i] for i in over180)
        p["opp"] = p["total"] < -immaterial          # credit bal for debtors / debit for creditors
        flags = []
        if p["opp"]: flags.append("Advance / opposite balance")
        if p["total"] >= immaterial and p["over180"] / p["total"] > 0.5: flags.append(">50% older than 180 days")
        if re.search(r"provision|ecl|doubtful", p["name"], re.I): flags.append("Provision ledger - review classification")
        if re.search(r"\b(USD|EUR|GBP)\b", p["name"]): flags.append("Foreign-currency ledger")
        if any(x < -immaterial for x in p["buckets"]) and any(x > immaterial for x in p["buckets"]):
            flags.append("Unadjusted Dr & Cr in different buckets - knock off bill-wise in Tally")
        if not p["material"]: flags.append("Immaterial / rounding balance")
        p["flags"] = "; ".join(flags)
    mat = [p for p in d["parties"] if p["material"]]
    tot = sum(p["total"] for p in d["parties"])
    bt = [sum(p["buckets"][i] for p in d["parties"]) for i in range(n)]
    pos = sorted([p for p in mat if p["total"] > 0], key=lambda p: -p["total"])
    neg = [p for p in mat if p["opp"]]
    bpos = [sum(p["buckets"][i] for p in pos) for i in range(n)]   # ageing of parties owing / owed
    bneg = [sum(p["buckets"][i] for p in neg) for i in range(n)]
    gross = sum(p["total"] for p in pos)
    k = {"total": tot, "gross": gross, "advances": sum(p["total"] for p in neg),
         "over90": sum(bpos[i] for i in first_overdue), "over180": sum(bpos[i] for i in over180),
         "parties": len(mat), "open": len(pos), "opp_count": len(neg),
         "top5_share": (sum(p["total"] for p in pos[:5]) / gross) if gross else 0}
    checks = []
    if d["grand"]:
        g = d["grand"]
        if g["total"] is not None:
            checks.append(("Ledger total = Tally Grand Total", tot, g["total"]))
        for i, b in enumerate(B):
            checks.append((f"Bucket '{b}' = Tally Grand Total", bt[i], g["buckets"][i]))
    checks.append(("Sum of buckets = Pending Bills", sum(bt), tot))
    mov = None
    if prev:
        pm = {p["name"]: p for p in prev["parties"]}
        names = list(dict.fromkeys([p["name"] for p in d["parties"]] + list(pm)))
        cm = {p["name"]: p for p in d["parties"]}
        mov = []
        for nm in names:
            a = cm.get(nm, {}).get("total", 0.0); b = pm.get(nm, {}).get("total", 0.0)
            if abs(a) < immaterial and abs(b) < immaterial: continue
            st = "New" if nm not in pm or abs(b) < immaterial else ("Cleared" if nm not in cm or abs(a) < immaterial else "")
            mov.append({"name": nm, "prev": b, "curr": a, "chg": a - b, "status": st,
                        "prev90": pm.get(nm, {}).get("over90", 0.0) if nm in pm else 0.0,
                        "curr90": cm.get(nm, {}).get("over90", 0.0)})
        mov.sort(key=lambda m: -abs(m["chg"]))
        k["prev_total"] = sum(p["total"] for p in prev["parties"])
        k["prev_gross"] = prev["kpi"]["gross"]; k["prev_over90"] = prev["kpi"]["over90"]
        k["prev_label"] = prev["as_at"].strftime("%b-%y") if prev["as_at"] else "Prior"
    d.update(kpi=k, bucket_totals=bt, bucket_pos=bpos, bucket_neg=bneg, checks=checks, movement=mov, top=pos[:10],
             first_overdue=first_overdue, over180_idx=over180, oa=oa)
    return d

# ----------------------------------------------------------------- excel
thin = Side(style="thin", color="BFBFBF")
def _hdr(ws, r, c1, vals, fill=NAVY):
    for j, v in enumerate(vals):
        c = ws.cell(r, c1 + j, v)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=thin)

def _title(ws, d, sub):
    ws["A1"] = d["company"]; ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    lab = d["as_at"].strftime("%d %B %Y") if d["as_at"] else d["period"]
    ws["A2"] = f"{'Debtor' if d['kind']=='Debtors' else 'Creditor'} ageing - {sub} - as at {lab}"
    ws["A2"].font = Font(size=11, color="595959")
    ws["A3"] = f"Source: {d['file']} (Tally Group Outstandings). Amounts in base currency; " + \
               ("net = Debit - Credit." if d["kind"] == "Debtors" else "net = Credit - Debit.")
    ws["A3"].font = Font(size=9, italic=True, color="7F7F7F")

def write_excel(d, path):
    wb = openpyxl.Workbook()
    B = d["buckets"]; n = len(B); k = d["kpi"]
    # ---------- Ageing Detail (built first; Dashboard formulas point here)
    det = wb.active; det.title = "Ageing Detail"
    _title(det, d, "party-wise")
    cols = ["Code", "Party", "Group"] + B + ["Total", "% of total", "Over 90 days", "Oldest bucket", "Flags"]
    _hdr(det, 5, 1, cols); det.row_dimensions[5].height = 32
    r0 = 6
    rows = sorted([p for p in d["parties"] if p["material"]], key=lambda p: -p["total"])
    fb, lb = 4, 3 + n                                    # first/last bucket col
    tc = lb + 1; last = r0 + len(rows) - 1; totr = last + 1
    for i, p in enumerate(rows):
        r = r0 + i
        det.cell(r, 1, int(p["code"]) if p["code"] else None); det.cell(r, 2, p["name"]); det.cell(r, 3, p["group"])
        for j, v in enumerate(p["buckets"]):
            det.cell(r, fb + j, round(v, 2)).number_format = NUM
        det.cell(r, tc, f"=SUM({CL(fb)}{r}:{CL(lb)}{r})").number_format = NUM
        det.cell(r, tc + 1, f"=IF(${CL(tc)}${totr}=0,0,{CL(tc)}{r}/${CL(tc)}${totr})").number_format = PCT
        ov = "+".join(f"{CL(fb+j)}{r}" for j in d["first_overdue"]) or "0"
        det.cell(r, tc + 2, f"={ov}").number_format = NUM
        det.cell(r, tc + 3, p["oldest"]); det.cell(r, tc + 4, p["flags"])
        if i % 2:
            for c in range(1, tc + 5): det.cell(r, c).fill = PatternFill("solid", fgColor=GREY)
    det.cell(totr, 2, "Total").font = Font(bold=True)
    for c in range(fb, tc + 3):
        if c == tc + 1: continue
        cell = det.cell(totr, c, f"=SUM({CL(c)}{r0}:{CL(c)}{last})")
        cell.number_format = NUM; cell.font = Font(bold=True); cell.border = Border(top=thin, bottom=Side(style="double"))
    det.cell(totr, tc + 1, f"=IF({CL(tc)}{totr}=0,0,1)").number_format = PCT
    det.conditional_formatting.add(f"{CL(fb)}{r0}:{CL(tc)}{last}", CellIsRule(operator="lessThan", formula=["-0.5"], font=Font(color="C00000")))
    det.freeze_panes = "D6"; det.auto_filter.ref = f"A5:{CL(tc+4)}{last}"
    widths = [8, 46, 26] + [15] * n + [16, 10, 15, 16, 44]
    for i, w in enumerate(widths): det.column_dimensions[CL(i + 1)].width = w
    D = "'Ageing Detail'!"

    # ---------- Dashboard
    db = wb.create_sheet("Dashboard", 0)
    _title(db, d, "dashboard")
    db.sheet_view.showGridLines = False
    side = ("Receivables", "Advances from customers") if d["kind"] == "Debtors" else ("Payables", "Advances to suppliers")
    TR = f"{D}${CL(tc)}${r0}:${CL(tc)}${last}"           # party total column in Ageing Detail
    tiles = [(f"Gross {side[0].lower()} (parties with balance due)", f'=SUMIF({TR},">0")', NUM),
             (side[1], f'=SUMIF({TR},"<0")', NUM),
             ("Net balance (agrees to Tally)", f"={D}{CL(tc)}{totr}", NUM),
             ("Overdue > 90 days (of gross)", None, NUM),
             ("% of gross > 90 days", "=IF(B6=0,0,B9/B6)", PCT),
             ("Older than 180 days (of gross)", None, NUM),
             ("Parties with balance due / in advance", f"{k['open']} / {k['opp_count']}", None),
             ("Top-5 concentration (of gross)", k["top5_share"], PCT)]
    if "prev_total" in k:
        tiles += [(f"Gross {side[0].lower()} {k['prev_label']}", round(k["prev_gross"], 2), NUM),
                  ("Change in gross vs prior month", "=B6-B14", NUM),
                  (f"Overdue > 90 days {k['prev_label']}", round(k["prev_over90"], 2), NUM)]
    _hdr(db, 5, 1, ["Key indicator", "Value"])
    br = 6 + len(tiles) + 2                              # bucket table header row
    ov_rows = [br + 1 + i for i in d["first_overdue"]]; o180 = [br + 1 + i for i in d["over180_idx"]]
    tiles[3] = (tiles[3][0], "=" + ("+".join(f"B{r}" for r in ov_rows) or "0"), NUM)
    tiles[5] = (tiles[5][0], "=" + ("+".join(f"B{r}" for r in o180) or "0"), NUM)
    for i, (a_, v, f) in enumerate(tiles):
        db.cell(6 + i, 1, a_); c = db.cell(6 + i, 2, v); c.font = Font(bold=True, size=12)
        if f: c.number_format = f
        else: c.alignment = Alignment(horizontal="right")
    _hdr(db, br, 1, ["Ageing bucket", side[0], "% of gross", side[1], "Nil-net parties (unadjusted)", "Net (per Tally)"])
    for j, b in enumerate(B):
        r = br + 1 + j; col = f"{D}${CL(fb+j)}${r0}:${CL(fb+j)}${last}"
        db.cell(r, 1, b)
        db.cell(r, 2, f'=SUMIF({TR},">0",{col})').number_format = NUM
        db.cell(r, 3, f"=IF($B${br+n+1}=0,0,B{r}/$B${br+n+1})").number_format = PCT
        db.cell(r, 4, f'=SUMIF({TR},"<0",{col})').number_format = NUM
        db.cell(r, 6, f"={D}{CL(fb+j)}{totr}").number_format = NUM
        db.cell(r, 5, f"=F{r}-B{r}-D{r}").number_format = NUM
    r = br + n + 1
    db.cell(r, 1, "Total").font = Font(bold=True)
    for c_, f in zip(range(2, 7), [NUM, PCT, NUM, NUM, NUM]):
        cc = db.cell(r, c_, f"=SUM({CL(c_)}{br+1}:{CL(c_)}{br+n})"); cc.number_format = f; cc.font = Font(bold=True)
        cc.border = Border(top=thin, bottom=Side(style="double"))
    ch = BarChart(); ch.type = "col"; ch.title = f"Ageing of gross {side[0].lower()}"; ch.style = 2
    ch.add_data(Reference(db, min_col=2, min_row=br, max_row=br + n), titles_from_data=True)
    ch.set_categories(Reference(db, min_col=1, min_row=br + 1, max_row=br + n))
    ch.legend = None; ch.height = 7.5; ch.width = 16; ch.y_axis.numFmt = "#,##0"; ch.y_axis.majorGridlines = None
    ch.series[0].graphicalProperties.solidFill = BLUE
    ch.dataLabels = DataLabelList(); ch.dataLabels.showVal = True; ch.dataLabels.numFmt = '#,##0.0,,"M"'
    db.add_chart(ch, "H5")
    # top 10
    tr = r + 3
    _hdr(db, tr, 1, ["Top 10 exposures", "Total", "% of gross", "Over 90 days"])
    top = d["top"]
    for i, p in enumerate(top):
        rr = tr + 1 + i
        db.cell(rr, 1, p["name"]); db.cell(rr, 2, round(p["total"], 2)).number_format = NUM
        db.cell(rr, 3, p["total"] / k["gross"] if k["gross"] else 0).number_format = PCT
        db.cell(rr, 4, round(p["over90"], 2)).number_format = NUM
    ch2 = BarChart(); ch2.type = "bar"; ch2.title = "Top exposures"; ch2.style = 2
    ch2.add_data(Reference(db, min_col=2, min_row=tr, max_row=tr + len(top)), titles_from_data=True)
    ch2.set_categories(Reference(db, min_col=1, min_row=tr + 1, max_row=tr + len(top)))
    ch2.legend = None; ch2.height = 8.5; ch2.width = 16; ch2.x_axis.scaling.orientation = "maxMin"
    ch2.y_axis.numFmt = "#,##0"; ch2.y_axis.majorGridlines = None
    ch2.series[0].graphicalProperties.solidFill = "1C5CAB"
    db.add_chart(ch2, "H21")
    for col, w in zip("ABCDEF", [44, 18, 12, 20, 18, 16]): db.column_dimensions[col].width = w

    # ---------- Advances / opposite balances
    adv = wb.create_sheet("Advances" if d["kind"] == "Debtors" else "Debit balances")
    _title(adv, d, "parties with opposite (advance) balances")
    _hdr(adv, 5, 1, ["Party", "Net balance"] + B)
    ar = 6
    for p in sorted([p for p in d["parties"] if p["material"] and p["opp"]], key=lambda p: p["total"]):
        adv.cell(ar, 1, p["name"]); adv.cell(ar, 2, round(p["total"], 2)).number_format = NUM
        for j, v in enumerate(p["buckets"]): adv.cell(ar, 3 + j, round(v, 2)).number_format = NUM
        ar += 1
    if ar == 6: adv.cell(6, 1, "None this month")
    adv.column_dimensions["A"].width = 46
    for c in range(2, 3 + n): adv.column_dimensions[CL(c)].width = 15

    # ---------- Movement
    if d["movement"] is not None:
        mv = wb.create_sheet("Movement")
        _title(mv, d, f"movement vs {d['kpi']['prev_label']}")
        _hdr(mv, 5, 1, ["Party", f"Total {d['kpi']['prev_label']}", "Total current", "Change", "Over 90 prior", "Over 90 current", "Status"])
        for i, m in enumerate(d["movement"]):
            r = 6 + i
            mv.cell(r, 1, m["name"])
            for j, key in enumerate(["prev", "curr"]): mv.cell(r, 2 + j, round(m[key], 2)).number_format = NUM
            mv.cell(r, 4, f"=C{r}-B{r}").number_format = NUM
            mv.cell(r, 5, round(m["prev90"], 2)).number_format = NUM; mv.cell(r, 6, round(m["curr90"], 2)).number_format = NUM
            mv.cell(r, 7, m["status"])
        e = 6 + len(d["movement"])
        mv.cell(e, 1, "Total").font = Font(bold=True)
        for c in range(2, 7):
            mv.cell(e, c, f"=SUM({CL(c)}6:{CL(c)}{e-1})").number_format = NUM; mv.cell(e, c).font = Font(bold=True)
        mv.column_dimensions["A"].width = 46
        for c in "BCDEF": mv.column_dimensions[c].width = 16

    # ---------- Checks
    ck = wb.create_sheet("Checks")
    _title(ck, d, "reconciliation & exceptions")
    _hdr(ck, 5, 1, ["Check", "Computed", "Per Tally", "Difference", "Result"])
    for i, (a, x, y) in enumerate(d["checks"]):
        r = 6 + i
        ck.cell(r, 1, a); ck.cell(r, 2, round(x, 2)).number_format = NUM; ck.cell(r, 3, round(y, 2)).number_format = NUM
        ck.cell(r, 4, f"=B{r}-C{r}").number_format = '#,##0.00;(#,##0.00);"-"'
        ck.cell(r, 5, f'=IF(ABS(D{r})<1,"OK","DIFFERENCE")')
    ck.conditional_formatting.add(f"E6:E{5+len(d['checks'])}", CellIsRule(operator="equal", formula=['"DIFFERENCE"'], font=Font(bold=True, color="C00000")))
    er = 8 + len(d["checks"])
    _hdr(ck, er, 1, ["Exception", "Party", "Net balance"])
    ex = [p for p in d["parties"] if p["flags"]]
    for i, p in enumerate(ex):
        ck.cell(er + 1 + i, 1, p["flags"]); ck.cell(er + 1 + i, 2, p["name"])
        ck.cell(er + 1 + i, 3, round(p["total"], 4)).number_format = '#,##0.00;(#,##0.00);"-"'
    ck.column_dimensions["A"].width = 52; ck.column_dimensions["B"].width = 46
    for c in "CDE": ck.column_dimensions[c].width = 16

    # ---------- Source (copy of raw Tally values)
    src = wb.create_sheet("Source")
    for row in d["ws"].iter_rows(values_only=True): src.append(list(row))
    src.column_dimensions["A"].width = 46
    wb.save(path)

# ----------------------------------------------------------------- html
def write_html(reports, path):
    def pack(d):
        k = d["kpi"]
        return {"kind": d["kind"], "company": d["company"],
                "asAt": d["as_at"].strftime("%d %b %Y") if d["as_at"] else d["period"],
                "buckets": d["buckets"], "bucketTotals": d["bucket_totals"], "bucketPos": d["bucket_pos"], "bucketNeg": d["bucket_neg"], "kpi": k,
                "parties": [{k2: p[k2] for k2 in ("name", "group", "buckets", "total", "over90", "oldest", "flags")}
                            for p in sorted(d["parties"], key=lambda p: -p["total"]) if p["material"]],
                "checks": [{"name": a, "diff": x - y} for a, x, y in d["checks"]],
                "movement": d["movement"][:15] if d["movement"] else None}
    data = json.dumps([pack(d) for d in reports], default=str)
    tpl = open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "dashboard_template.html"), encoding="utf-8").read()
    open(path, "w", encoding="utf-8").write(tpl.replace("/*__DATA__*/[]", data))

# ----------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("inputs", nargs="+"); ap.add_argument("--prev", nargs="*", default=[])
    ap.add_argument("--out", default="."); ap.add_argument("--immaterial", type=float, default=1.0)
    a = ap.parse_args(); os.makedirs(a.out, exist_ok=True)
    prevs = {}
    for f in a.prev:
        p = analyse(parse(f), immaterial=a.immaterial); prevs[p["kind"]] = p
    reports = []
    for f in a.inputs:
        d = parse(f); analyse(d, prevs.get(d["kind"]), a.immaterial)
        tag = d["as_at"].strftime("%b-%y") if d["as_at"] else "period"
        out = os.path.join(a.out, f"{d['kind'][:-1]}_Ageing_{tag}.xlsx")
        write_excel(d, out); reports.append(d)
        bad = [c for c in d["checks"] if abs(c[1] - c[2]) >= 1]
        print(f"{d['kind']}: {len(d['parties'])} ledgers, net {d['kpi']['total']:,.2f}, "
              f"checks {'OK' if not bad else 'DIFFERENCES: ' + str(len(bad))} -> {out}")
    tag = reports[0]["as_at"].strftime("%b-%y") if reports[0]["as_at"] else "period"
    hp = os.path.join(a.out, f"Ageing_Dashboard_{tag}.html"); write_html(reports, hp); print("Dashboard ->", hp)

if __name__ == "__main__":
    main()
