"""
Writes the management ageing workbook in the house format:
  Debtors Dashboard | Creditor Dashboard | DR_Row Data | CR_Row Data | FX Rates | Checks

Each dashboard: title, period, four KPI tiles, a one-line commentary, then two tables -
balances due (receivables / payables) and advances (customer advances / advances paid),
with ageing buckets and the Pending Bills amount converted to each foreign currency at the
rate on the FX Rates sheet (formulas, so a changed rate flows through).
"""
import re
from copy import copy
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL

FONT = "Calibri"
THIN = Side(style="thin", color="8EA9C1")
DOUBLE = Side(style="double", color="1F4E78")
fill = lambda c: PatternFill("solid", fgColor=c)

STYLE = {
    "Debtors": dict(title="DEBTORS REPORT", kpi=["TOTAL RECEIVABLES", "CUSTOMER ADVANCES", "NET RECEIVABLE", ">90 DAYS RECEIVABLE"],
                    s1="RECEIVABLES (NET DEBIT BALANCES)", t1="TOTAL RECEIVABLES",
                    s2="CUSTOMER ADVANCES (NET CREDIT BALANCES)", t2="TOTAL CUSTOMER ADVANCES",
                    other="OTHER CUSTOMERS", sheet="Debtors Dashboard", raw="DR_Row Data",
                    row_fmt='_(* #,##0_);_(* \\(#,##0\\);_(* "-"??_);_(@_)', tot_fmt='#,##0;[Red]\\(#,##0\\);\\-',
                    kpi_fmt='#,##0.00;[Red]\\(#,##0.00\\);\\-', kpi_fill="EAF3F8", hdr_fill="D9E2F3", tot_fill="D9EAF7", row_size=11),
    "Creditors": dict(title="CREDITOR REPORT", kpi=["TOTAL PAYABLES", "ADVANCES PAID", "NET PAYABLE", ">90 DAYS PAYABLE"],
                      s1="PAYABLES (NET CREDIT BALANCES)", t1="TOTAL PAYABLES",
                      s2="ADVANCES PAID (DEBIT BALANCES)", t2="TOTAL ADVANCES PAID",
                      other="OTHER SUPPLIERS", sheet="Creditor Dashboard", raw="CR_Row Data",
                      row_fmt='#,##0.00;[Red]\\(#,##0.00\\);\\-', tot_fmt='#,##0.00;[Red]\\(#,##0.00\\);\\-',
                      kpi_fmt='#,##0;[Red]\\(#,##0\\);\\-', kpi_fill="D9EAF7", hdr_fill="D9EAF1", tot_fill="C8DEE8", row_size=10),
}
CUR_FMT = {"USD": '_("$"* #,##0_);_("$"* \\(#,##0\\);_("$"* "-"??_);_(@_)',
           "EUR": '_("€"* #,##0_);_("€"* \\(#,##0\\);_("€"* "-"??_);_(@_)',
           "GBP": '_("£"* #,##0_);_("£"* \\(#,##0\\);_("£"* "-"??_);_(@_)'}

def bucket_label(b):
    s = re.sub(r"[()]", "", b).strip()
    s = re.sub(r"\s+to\s+", "-", s)
    return re.sub(r"\s+", " ", s)

def _sections(d, immaterial):
    """Split parties into (balances due, advances) keeping Tally order and sub-groups.
    Section values are shown positive: advances are sign-flipped."""
    top = "Sundry " + d["kind"]
    has_groups = any(p["group"] != top for p in d["parties"])
    due, adv = [], []
    for p in d["parties"]:
        if abs(p["total"]) < immaterial and all(abs(x) < immaterial for x in p["buckets"]):
            continue
        grp = None
        if has_groups:
            grp = STYLE[d["kind"]]["other"] if p["group"] == top else p["group"].upper()
        row = {"name": p["name"], "group": grp}
        bk = [0.0 if abs(x) < immaterial else x for x in p["buckets"]]   # hide rounding residues, as Tally users do
        if p["total"] >= 0:
            due.append(dict(row, total=p["total"], buckets=bk))
        else:
            adv.append(dict(row, total=-p["total"], buckets=[-x for x in bk]))
    def order(rows):          # group blocks in order of first appearance, ledgers in Tally order
        if not has_groups: return rows
        seen = []
        for r in rows:
            if r["group"] not in seen: seen.append(r["group"])
        if STYLE[d["kind"]]["other"] in seen:               # direct ledgers go last
            seen.remove(STYLE[d["kind"]]["other"]); seen.append(STYLE[d["kind"]]["other"])
        return [r for g in seen for r in rows if r["group"] == g]
    return order(due), order(adv), has_groups

def _box(ws, r1, r2, c1, c2):
    for r in range(r1, r2 + 1):
        for c in (c1, c2):
            cell = ws.cell(r, c)
            b = copy(cell.border)
            cell.border = Border(left=THIN if c == c1 else b.left, right=THIN if c == c2 else b.right,
                                 top=b.top, bottom=b.bottom)

def dashboard(wb, d, fx_rows, immaterial=1.0):
    st = STYLE[d["kind"]]
    ws = wb.create_sheet(st["sheet"])
    ws.sheet_view.showGridLines = False
    B = d["buckets"]; n = len(B)
    curs = list(fx_rows)                                   # [(code, fx sheet cell ref)]
    last = 2 + n + len(curs)                               # last column index
    LC = CL(last)
    due, adv, grouped = _sections(d, immaterial)

    # numbers for the commentary (formulas carry the live values)
    g1 = sum(r["total"] for r in due); g2 = sum(r["total"] for r in adv)
    oa = next((i for i, b in enumerate(B) if "account" in b.lower()), None)
    aged = [i for i in range(n) if i != oa]
    over = aged[1:]
    o90 = sum(r["buckets"][i] for r in due for i in over)

    # title + period
    ws.merge_cells(f"A1:{LC}1"); c = ws["A1"]
    c.value = f"{d['company'].upper()} — {st['title']}"
    c.font = Font(FONT, 18, bold=True, color="FFFFFF"); c.fill = fill("17365D")
    c.alignment = Alignment("center", "center"); ws.row_dimensions[1].height = 32.1
    ws.merge_cells(f"A2:{LC}2"); c = ws["A2"]
    c.value = f"Period: {d['period']}"; c.font = Font(FONT, 11, bold=True); c.alignment = Alignment("center")
    if curs:
        ws.merge_cells(f"A3:{LC}3"); c = ws["A3"]
        c.value = "Foreign currency at " + "; ".join(f"1 {k} = {v[1]:,.4f} AED ({v[2]})" for k, v in fx_rows.items())
        c.font = Font(FONT, 9, italic=True, color="595959"); c.alignment = Alignment("center")

    # KPI tiles (pairs of columns A:B, C:D, E:F, G:H)
    for j, lab in enumerate(st["kpi"]):
        c1 = 1 + 2 * j
        ws.merge_cells(start_row=4, start_column=c1, end_row=4, end_column=c1 + 1)
        ws.merge_cells(start_row=5, start_column=c1, end_row=5, end_column=c1 + 1)
        h = ws.cell(4, c1, lab); h.font = Font(FONT, 11, bold=True, color="FFFFFF"); h.fill = fill("1F4E78")
        h.alignment = Alignment("center", "center")
        v = ws.cell(5, c1); v.font = Font(FONT, 16, bold=True, color="1F1F1F"); v.fill = fill(st["kpi_fill"])
        v.alignment = Alignment("center", "center"); v.number_format = st["kpi_fmt"]
        for cc in (c1, c1 + 1):
            ws.cell(4, cc).border = Border(top=THIN, left=THIN if cc == 1 else None, right=THIN if cc == 8 else None)
            ws.cell(5, cc).border = Border(bottom=THIN, left=THIN if cc == 1 else None, right=THIN if cc == 8 else None)
            ws.cell(4, cc).fill = fill("1F4E78"); ws.cell(5, cc).fill = fill(st["kpi_fill"])
    ws.row_dimensions[5].height = 27.75

    # commentary
    ws.merge_cells(f"A7:{LC}8"); c = ws["A7"]
    if d["kind"] == "Debtors":
        c.value = (f"Gross receivables are {g1:,.0f}, while customer advances received are {g2:,.0f}. "
                   f"Net receivables are {g1 - g2:,.0f}. Receivables aged more than 90 days are {o90:,.0f}.")
    else:
        c.value = (f"Gross payables are {g1:,.0f}, while advances paid are {g2:,.0f}. Net payable is {g1 - g2:,.0f}. "
                   f"Payables aged more than 90 days are {o90:,.0f}." + (" Supplier balances are presented by source group." if grouped else ""))
    c.font = Font(FONT, 11); c.fill = fill("F3F6F9"); c.alignment = Alignment("center", "center", wrap_text=True)
    for r in (7, 8):
        for cc in range(1, last + 1):
            ws.cell(r, cc).border = Border(top=THIN if r == 7 else None, bottom=THIN if r == 8 else None,
                                           left=THIN if cc == 1 else None, right=THIN if cc == last else None)

    def table(r, title, rows, total_label):
        ws.merge_cells(f"A{r}:{LC}{r}"); t = ws.cell(r, 1, title)
        t.font = Font(FONT, 12, bold=True, color="FFFFFF"); t.fill = fill("0F6B78"); t.alignment = Alignment("center", "center")
        ws.row_dimensions[r].height = 24
        r += 1
        heads = ["Particulars", "Pending Bills"] + [bucket_label(b) for b in B] + [f"Amount in {k}" for k in curs]
        for j, h in enumerate(heads):
            c = ws.cell(r, 1 + j, h); c.font = Font(FONT, 11, bold=True); c.fill = fill(st["hdr_fill"])
            c.alignment = Alignment("center", "center", wrap_text=True); c.border = Border(bottom=THIN)
        hdr_row = r; r += 1; first = r; grp = None; band = 0
        for row in rows:
            if row["group"] and row["group"] != grp:
                grp = row["group"]
                for cc in range(1, last + 1): ws.cell(r, cc).fill = fill("EAF2F5")
                g = ws.cell(r, 1, grp); g.font = Font(FONT, 11, bold=True); g.alignment = Alignment("left", "center")
                r += 1; band = 0
            a = ws.cell(r, 1, row["name"]); a.number_format = "@"
            a.font = Font(FONT, st["row_size"]); a.alignment = Alignment("left", "center", indent=2 if row["group"] else 0)
            vals = [row["total"]] + row["buckets"]
            for j, v in enumerate(vals):
                c = ws.cell(r, 2 + j, round(v, 4)); c.number_format = st["row_fmt"]
                c.font = Font(FONT, st["row_size"]); c.alignment = Alignment("right", "center")
            for j, (k, ref) in enumerate(fx_rows.items()):
                c = ws.cell(r, 3 + n + j, f"=B{r}/{ref[0]}"); c.number_format = CUR_FMT.get(k, st["row_fmt"])
                c.font = Font(FONT, st["row_size"]); c.alignment = Alignment("right", "center")
            if band % 2 == 0:
                for cc in range(1, last + 1): ws.cell(r, cc).fill = fill("F2F2F2")
            band += 1; r += 1
        lastrow = r - 1
        t = ws.cell(r, 1, total_label); t.font = Font(FONT, 11, bold=True); t.alignment = Alignment("center", "center")
        for cc in range(1, last + 1):
            c = ws.cell(r, cc); c.fill = fill(st["tot_fill"]); c.border = Border(top=DOUBLE, bottom=THIN)
            if cc >= 2:
                c.value = f"=SUM({CL(cc)}{first}:{CL(cc)}{lastrow})" if lastrow >= first else 0
                c.font = Font(FONT, 11, bold=True, color="16324F"); c.alignment = Alignment("right", "center")
                c.number_format = CUR_FMT.get(curs[cc - 3 - n], st["tot_fmt"]) if cc > 2 + n else st["tot_fmt"]
        _box(ws, hdr_row - 1, r, 1, last)
        return r

    tot1 = table(10, st["s1"], due, st["t1"])
    s2 = tot1 + 3
    tot2 = table(s2, st["s2"], adv, st["t2"])

    ws["A5"] = f"=B{tot1}"; ws["C5"] = f"=B{tot2}"; ws["E5"] = "=A5-C5"
    ws["G5"] = "=" + ("+".join(f"{CL(3 + i)}{tot1}" for i in over) or "0")

    ws.column_dimensions["A"].width = 57.14
    for cc in range(2, 3 + n): ws.column_dimensions[CL(cc)].width = 15.57
    for cc in range(3 + n, last + 1): ws.column_dimensions[CL(cc)].width = 18.14
    ws.freeze_panes = None
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return {"due": due, "adv": adv, "g1": g1, "g2": g2, "o90": o90}

def raw_copy(wb, d, title):
    """Copy the Tally sheet as-is (values, number formats, bold/italic/indent, merges, widths)."""
    src = d["ws"]; ws = wb.create_sheet(title)
    for row in src.iter_rows():
        for c in row:
            if c.value is None and not c.has_style: continue
            t = ws.cell(c.row, c.column, c.value)
            if c.has_style:
                t.number_format = c.number_format
                t.font = Font(name=c.font.name, size=c.font.sz, bold=c.font.b, italic=c.font.i)
                t.alignment = Alignment(horizontal=c.alignment.horizontal, vertical=c.alignment.vertical,
                                        indent=c.alignment.indent or 0, wrap_text=c.alignment.wrap_text)
    for m in src.merged_cells.ranges: ws.merge_cells(str(m))
    for k, dim in src.column_dimensions.items():
        if dim.width: ws.column_dimensions[k].width = dim.width
    if not src.column_dimensions["A"].width: ws.column_dimensions["A"].width = 52

def fx_sheet(wb, fx):
    ws = wb.create_sheet("FX Rates")
    ws["A1"] = "Foreign exchange rates - AED per 1 unit of currency"; ws["A1"].font = Font(FONT, 14, bold=True, color="17365D")
    ws["A2"] = ("Yellow cells are inputs: change a rate and every converted amount in the dashboards updates. "
                "USD is pegged by the Central Bank of the UAE at 3.6725 AED.")
    ws["A2"].font = Font(FONT, 9, italic=True, color="595959")
    for j, h in enumerate(["Currency", "AED per unit", "Source", "Rate date"]):
        c = ws.cell(4, 1 + j, h); c.font = Font(FONT, 11, bold=True, color="FFFFFF"); c.fill = fill("1F4E78")
        c.alignment = Alignment("center")
    refs = {}
    for i, (k, v) in enumerate(fx.items()):
        r = 5 + i
        ws.cell(r, 1, k).font = Font(FONT, 11, bold=True)
        c = ws.cell(r, 2, v["rate"]); c.number_format = "0.0000"; c.fill = fill("FFF2CC")
        ws.cell(r, 3, v["source"]); ws.cell(r, 4, v["date"])
        refs[k] = (f"'FX Rates'!$B${r}", v["rate"], v["source"])
    for col, w in zip("ABCD", [12, 16, 40, 14]): ws.column_dimensions[col].width = w
    return refs

def checks_sheet(wb, reports):
    ws = wb.create_sheet("Checks")
    ws["A1"] = "Reconciliation to Tally Grand Total and exceptions"; ws["A1"].font = Font(FONT, 14, bold=True, color="17365D")
    ws["A2"] = ("A DIFFERENCE means the ledger rows in the export don't add up to Tally's Grand Total - "
                "usually rows were deleted or edited after export. Re-export from Tally rather than editing.")
    ws["A2"].font = Font(FONT, 9, italic=True, color="595959")
    r = 4
    for d in reports:
        ws.cell(r, 1, f"{d['kind']} - {d['file']} [{d['sheet']}]").font = Font(FONT, 12, bold=True); r += 1
        for j, h in enumerate(["Check", "Ledgers", "Per Tally", "Difference", "Result"]):
            c = ws.cell(r, 1 + j, h); c.font = Font(FONT, 11, bold=True, color="FFFFFF"); c.fill = fill("1F4E78")
        r += 1
        for name, x, y in d["checks"]:
            ws.cell(r, 1, name)
            for j, v in enumerate([x, y]): ws.cell(r, 2 + j, round(v, 2)).number_format = '#,##0.00;[Red]\\(#,##0.00\\);\\-'
            ws.cell(r, 4, f"=B{r}-C{r}").number_format = '#,##0.00;[Red]\\(#,##0.00\\);\\-'
            res = ws.cell(r, 5, f'=IF(ABS(D{r})<1,"OK","DIFFERENCE")'); res.font = Font(FONT, 11, bold=True)
            r += 1
        flagged = [p for p in d["parties"] if p.get("flags")]
        if flagged:
            r += 1
            for j, h in enumerate(["Exception", "Party", "Net balance"]):
                c = ws.cell(r, 1 + j, h); c.font = Font(FONT, 11, bold=True); c.fill = fill("D9E2F3")
            r += 1
            for p in flagged:
                ws.cell(r, 1, p["flags"]); ws.cell(r, 2, p["name"])
                ws.cell(r, 3, round(p["total"], 2)).number_format = '#,##0.00;[Red]\\(#,##0.00\\);\\-'; r += 1
        r += 2
    from openpyxl.formatting.rule import CellIsRule
    ws.conditional_formatting.add(f"E4:E{r}", CellIsRule(operator="equal", formula=['"DIFFERENCE"'], font=Font(bold=True, color="C00000")))
    for col, w in zip("ABCDE", [62, 48, 18, 16, 14]): ws.column_dimensions[col].width = w

def build(reports, fx, path, immaterial=1.0):
    """reports: analysed current-period reports (at most one per kind)."""
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    by = {d["kind"]: d for d in reports}
    refs = fx_sheet(wb, fx)
    summary = {}
    for kind in ("Debtors", "Creditors"):
        if kind in by: summary[kind] = dashboard(wb, by[kind], refs, immaterial)
    for kind in ("Debtors", "Creditors"):
        if kind in by:
            raw_copy(wb, by[kind], STYLE[kind]["raw"])
            wb[STYLE[kind]["raw"]].sheet_state = "hidden"          # as in the house file; unhide to view
    wb.move_sheet("FX Rates", offset=len(wb.sheetnames))
    checks_sheet(wb, [by[k] for k in ("Debtors", "Creditors") if k in by])
    wb.active = 0
    wb.save(path)
    return summary
